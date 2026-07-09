#!/usr/bin/env python3
"""Generate Livy delegation token renewal documentation (DOCX) – RPC-only flow."""

import os
from datetime import date

from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from PIL import Image, ImageDraw, ImageFont

OUTPUT = os.path.join(os.path.dirname(__file__), "livy-delegation-token-renewal.docx")
DIAGRAM_DIR = os.path.join(os.path.dirname(__file__), "diagrams")


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    return p


def add_bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def add_numbered(doc, text):
    doc.add_paragraph(text, style="List Number")


def add_code_block(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    p.paragraph_format.left_indent = Cm(0.5)
    return p


def _font(size=14):
    for name in ("Arial.ttf", "DejaVuSans.ttf", "Helvetica.ttc"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def _box(draw, xy, text, fill, font, outline="#333333"):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=10, fill=fill, outline=outline, width=2)
    lines = text.split("\n")
    line_h = font.size + 4
    total_h = len(lines) * line_h
    y = y1 + ((y2 - y1) - total_h) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        tw = bbox[2] - bbox[0]
        draw.text((x1 + (x2 - x1 - tw) / 2, y), line, fill="#1a1a1a", font=font)
        y += line_h


def _arrow(draw, start, end, color="#555555"):
    draw.line([start, end], fill=color, width=2)
    ex, ey = end
    sx, sy = start
    import math
    angle = math.atan2(ey - sy, ex - sx)
    size = 10
    p1 = (ex - size * math.cos(angle - 0.4), ey - size * math.sin(angle - 0.4))
    p2 = (ex - size * math.cos(angle + 0.4), ey - size * math.sin(angle + 0.4))
    draw.polygon([end, p1, p2], fill=color)


def draw_lifecycle_diagram(path):
    w, h = 900, 520
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    font = _font(13)
    font_sm = _font(11)

    draw.text((20, 15), "Full lifecycle (RPC-only)", fill="#1a5276", font=_font(16))

    boxes = [
        (60, 60, 260, 120, "1. Client\nsession + proxyUser", "#D6E4F0"),
        (330, 60, 570, 120, "2. Livy server\nDelegationTokenManager.obtainTokens()", "#A9CCE3"),
        (640, 60, 860, 120, "3. spark-submit\n--proxy-user (initial tokens)", "#7FB3D5"),
        (60, 200, 300, 270, "4. SessionDelegationTokenRenewer\nbackground thread starts", "#D5F5E3"),
        (360, 200, 600, 270, "5. Renewal\nDelegationTokenManager.renewTokens()", "#ABEBC6"),
        (660, 200, 860, 270, "6. RPC push\nUpdateCredentialsRequest", "#82E0AA"),
        (200, 360, 500, 430, "Interactive: RSCClient.updateCredentials()", "#F9E79F"),
        (520, 360, 820, 430, "Batch: DelegationTokenRpcBootstrap\n+ DelegationTokenRpcRegistry", "#F5B041"),
        (300, 470, 600, 510, "Driver UGI.addCredentials() – no HDFS polling", "#E8DAEF"),
    ]
    for box in boxes:
        _box(draw, box[:4], box[4], box[5], font_sm if len(box[4]) > 30 else font)

    arrows = [
        ((260, 90), (330, 90)),
        ((570, 90), (640, 90)),
        ((160, 120), (160, 200)),
        ((480, 120), (480, 200)),
        ((750, 120), (750, 200)),
        ((300, 235), (360, 235)),
        ((600, 235), (660, 235)),
        ((430, 270), (350, 360)),
        ((760, 270), (700, 360)),
        ((350, 430), (450, 470)),
        ((700, 430), (550, 470)),
    ]
    for a, b in arrows:
        _arrow(draw, a, b)

    img.save(path)


def draw_renewal_sequence(path):
    w, h = 900, 480
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    font = _font(12)
    font_sm = _font(11)

    draw.text((20, 15), "Renewal cycle (interactive and batch)", fill="#1a5276", font=_font(16))

    actors = [
        (80, 50, "Client"),
        (280, 50, "Livy Server"),
        (500, 50, "Renewer thread"),
        (720, 50, "Spark Driver"),
    ]
    for x, y, name in actors:
        draw.text((x, y), name, fill="#1a1a1a", font=font)
        draw.line([(x + 40, y + 22), (x + 40, h - 30)], fill="#BBBBBB", width=1)

    steps = [
        (0, 90, 1, 2, "obtainTokens(proxyUser)"),
        (1, 130, 3, 1, "spark-submit --proxy-user"),
        (2, 170, 1, 2, "scheduleRenewal()"),
        (2, 220, 2, 2, "renewTokens() [real user]"),
        (2, 270, 2, 3, "serialize(credentials)"),
        (3, 310, 2, 3, "UpdateCredentialsRequest (RPC)"),
        (3, 350, 3, 3, "UGI.addCredentials()"),
        (2, 400, 2, 2, "session stop → renewer.stop()"),
    ]
    y_base = 70
    for src, y_off, src_idx, dst_idx, label in steps:
        y = y_base + y_off
        x1 = actors[src_idx][0] + 40
        x2 = actors[dst_idx][0] + 40
        _arrow(draw, (x1, y), (x2, y))
        draw.text(((x1 + x2) / 2 - len(label) * 3, y - 14), label, fill="#333333", font=font_sm)

    img.save(path)


def add_image(doc, path, caption, width=Inches(6.2)):
    if os.path.exists(path):
        doc.add_picture(path, width=width)
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(10)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = val
    doc.add_paragraph()


def build_document():
    os.makedirs(DIAGRAM_DIR, exist_ok=True)
    lifecycle_png = os.path.join(DIAGRAM_DIR, "dt-lifecycle-rpc.png")
    renewal_png = os.path.join(DIAGRAM_DIR, "dt-renewal-rpc.png")
    draw_lifecycle_diagram(lifecycle_png)
    draw_renewal_sequence(renewal_png)

    doc = Document()
    title = doc.add_heading("Livy – Delegation Token Renewal (RPC)", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph(f"Generated: {date.today().isoformat()}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_para(
        doc,
        "This document describes delegation token handling for Apache Livy proxy-user "
        "(impersonation) sessions. Renewed tokens are pushed to the Spark driver entirely "
        "over encrypted RSC RPC. There is no HDFS credentials file, no file polling, and "
        "no staging directory for tokens.")

    add_heading(doc, "1. Summary", 1)
    add_bullet(doc, "Activation: proxyUser + Kerberos + livy.impersonation.enabled + renewal.enabled")
    add_bullet(doc, "Initial tokens: Spark --proxy-user at launch")
    add_bullet(doc, "Renewal: Livy server (real user) → RPC push to driver")
    add_bullet(doc, "Supported services: hdfs, hive, hbase, kafka (+ ofs:// Ozone auto-discovery)")
    add_bullet(doc, "Provider registry: built-in + SPI + livy.impersonation.delegation-token.providers")
    add_bullet(doc, "Filesystem auto-discovery: hdfs://, ofs://, o3fs://, s3a, … from jars/files/conf")
    add_bullet(doc, "Session types: interactive and batch, client and cluster deploy mode")

    add_heading(doc, "2. Architecture", 1)
    add_image(doc, lifecycle_png, "Figure 1 – Full delegation token lifecycle (RPC-only)")

    add_heading(doc, "3. Launch flow", 1)
    add_numbered(doc, "Client creates a session with proxyUser (impersonation).")
    add_numbered(doc, "BatchSession / InteractiveSession calls SessionDelegationTokenRenewer.prepareForLaunch().")
    add_numbered(doc, "DelegationTokenManager.obtainTokens() acquires tokens as the Livy login user.")
    add_numbered(doc, "Spark starts with --proxy-user; initial delegation tokens arrive during launch.")
    add_numbered(doc, "In batch mode, DelegationTokenRpcRegistry sets spark conf (launcher address, client id/secret, extraListeners, spark.jars).")
    add_numbered(doc, "In interactive mode, token push uses the existing RSC connection.")

    add_heading(doc, "4. Renewal flow", 1)
    add_image(doc, renewal_png, "Figure 2 – Renewal cycle")

    add_numbered(doc, "SessionDelegationTokenRenewer schedules renewal on a background thread.")
    add_numbered(doc, "Interval: min(token_max_lifetime / 2, livy.impersonation.delegation-token.renewal.interval)")
    add_numbered(doc, "DelegationTokenManager.renewTokens() – real Livy user, Hadoop TokenRenewer SPI.")
    add_numbered(doc, "Credentials serialized → UpdateCredentialsRequest RPC message.")
    add_numbered(doc, "On the driver, DelegationTokenRpcBootstrap.applyCredentials() → UGI.addCredentials().")
    add_numbered(doc, "On session stop: renewer.stop() and RPC connections closed.")

    add_heading(doc, "5. Interactive vs. Batch", 1)

    add_table(doc, ["Aspect", "Interactive session", "Batch session"], [
        ("RPC channel", "Existing RSCClient.updateCredentials()", "DelegationTokenRpcRegistry ↔ DelegationTokenRpcBootstrap"),
        ("Driver listener", "RSCDriver (existing)", "spark.extraListeners → DelegationTokenRpcBootstrap"),
        ("Batch registry", "Not required", "Livy-side RpcServer waits for driver registration"),
        ("Deploy mode", "Client and cluster", "Client and cluster (K8s/cluster: spark.driver.host)"),
        ("HDFS polling", "None", "None"),
    ])

    add_heading(doc, "6. Main components", 1)

    add_table(doc, ["Component", "File", "Role"], [
        ("DelegationTokenManager", "server/.../DelegationTokenManager.scala",
         "Token obtain, renew, serialize, max-lifetime config lookup"),
        ("DelegationTokenProviderRegistry", "server/.../DelegationTokenProviderRegistry.scala",
         "Built-in + SPI + config provider resolution"),
        ("BuiltInDelegationTokenProviders", "server/.../BuiltInDelegationTokenProviders.scala",
         "Built-in hdfs, hive, hbase, kafka providers"),
        ("SessionFilesystemUriCollector", "server/.../SessionFilesystemUriCollector.scala",
         "Filesystem URI auto-discovery from session inputs"),
        ("SessionDelegationTokenRenewer", "server/.../SessionDelegationTokenRenewer.scala",
         "Per-session background thread, scheduled renewal and push"),
        ("DelegationTokenRpcRegistry", "server/.../DelegationTokenRpcRegistry.scala",
         "Batch driver registration and RPC push from Livy server"),
        ("DelegationTokenRpcBootstrap", "rsc/.../DelegationTokenRpcBootstrap.java",
         "Batch driver-side RPC server and Livy registration"),
        ("RSCDriver", "rsc/.../RSCDriver.java",
         "Interactive: receives UpdateCredentialsRequest"),
    ])

    add_heading(doc, "7. Configuration", 1)

    add_code_block(doc, """# Impersonation (prerequisite)
livy.impersonation.enabled = true

# Delegation token renewal (default: enabled)
livy.impersonation.delegation-token.renewal.enabled = true
livy.impersonation.delegation-token.renewal.interval = 1h
livy.impersonation.delegation-token.renewer.principal = livy/hostname@REALM
livy.impersonation.delegation-token.extra.filesystems = hdfs://ns2
livy.impersonation.delegation-token.services = hdfs,hive,hbase,kafka
livy.impersonation.delegation-token.auto-discover.filesystems = true
livy.impersonation.delegation-token.providers = mycloud=com.example.MyTokenProvider

# Set automatically in batch mode:
# spark.extraListeners = org.apache.livy.rsc.driver.DelegationTokenRpcBootstrap
# spark.__livy__.livy.rsc.launcher-address / launcher-port / client-id / client-secret
# spark.jars += livy-rsc.jar (if needed)""")

    add_para(doc, "Service-specific delegation token lifetime keys "
             "(DelegationTokenManager.getMaxLifetime):", bold=True)
    add_table(doc, ["Service", "Config key(s)", "IT mini cluster value"], [
        ("HDFS", "dfs.namenode.delegation.token.max-lifetime", "15s"),
        ("Hive", "hive.cluster.delegation.token.max-lifetime", "15s"),
        ("Hive", "hive.cluster.delegation.token.renew-interval", "8s"),
        ("Hive", "hive.cluster.delegation.token.gc-interval", "5s"),
        ("HBase", "hbase.auth.token.max.lifetime", "15s"),
        ("HBase", "hbase.auth.key.update.interval", "8s"),
        ("Kafka", "delegation.token.max.lifetime.ms", "15000"),
        ("Kafka", "delegation.token.expiry.time.ms", "7500"),
        ("Ozone", "ozone.manager.delegation.token.max-lifetime", "15s"),
        ("Ozone", "ozone.manager.delegation.token.renew-interval", "8s"),
        ("Ozone", "ozone.manager.delegation.remover.scan.interval", "5s"),
    ])

    add_para(doc, "Notes:", bold=True)
    add_bullet(doc, "Renewal is active only for proxy-user sessions (DelegationTokenManager.isEnabled).")
    add_bullet(doc, "HBase and Kafka tokens require the corresponding client libraries on the Livy classpath.")
    add_bullet(doc, "No livy.rsc.delegation.tokens.path and no poll-interval – purely RPC-based.")
    add_bullet(doc, "Ozone ofs:// URIs are auto-discovered via SessionFilesystemUriCollector.")

    add_heading(doc, "8. Test environment (integration-test mini cluster)", 1)

    add_para(
        doc,
        "The integration-test module starts an embedded KerberosMiniCluster by default "
        "(cluster.type=mini, kerberos.enabled=true). Tests run with Maven test-scope "
        "dependencies; no external HDFS/YARN/Livy installation is required.")

    add_heading(doc, "8.1. Mini cluster architecture", 2)

    add_table(doc, ["Layer", "Component", "Startup", "Maven artifact (test scope)"], [
        ("Base", "MiniDFSCluster + MiniYARNCluster + LivyServer", "subprocess (MiniHdfsMain, MiniYarnMain, MiniLivyMain)", "hadoop-hdfs, hadoop-yarn-server-tests"),
        ("Security", "MiniKdc (Kerberos)", "in-process", "hadoop-minikdc"),
        ("Embedded services", "MiniOzoneCluster", "in-process (reflection)", "ozone-mini-cluster 2.1.1"),
        ("Embedded services", "Hive Metastore (Derby + Thrift)", "in-process daemon thread", "hive-service 3.0.0"),
        ("Embedded services", "HBase mini cluster", "in-process (HBaseTestingUtility)", "hbase-server 2.4.17 (tests)"),
        ("Embedded services", "Kafka broker", "in-process (KafkaServer + Curator ZK)", "kafka_2.12 2.8.2, curator-test"),
    ])

    add_para(doc, "Startup order (KerberosMiniCluster.configureClusterConfig):", bold=True)
    add_numbered(doc, "MiniKdc + secure core-site.xml / hdfs-site.xml / yarn-site.xml (short token lifetimes).")
    add_numbered(doc, "MiniOzoneCluster (ozone.enabled=true by default).")
    add_numbered(doc, "Hive Metastore (hive.enabled=true).")
    add_numbered(doc, "HBase mini cluster (hbase.enabled=true).")
    add_numbered(doc, "Kafka broker (kafka.enabled=true).")
    add_numbered(doc, "HDFS / YARN / Livy subprocesses (shared HADOOP_CONF_DIR).")

    add_heading(doc, "8.2. KDC principals and proxy user", 2)
    add_bullet(doc, "Principals: hdfs, HTTP, yarn, livy, proxy, om, scm, hive, hbase, kafka")
    add_bullet(doc, "Livy login: livy/localhost@REALM (keytab)")
    add_bullet(doc, "Proxy user for tests: proxy/localhost@REALM → livy.test.proxyUser=proxy")
    add_bullet(doc, "hadoop.proxyuser.livy.users=* and hosts=* in core-site.xml")

    add_heading(doc, "8.3. Short token lifetimes (renewal IT)", 2)
    add_table(doc, ["Setting", "Value", "Purpose"], [
        ("Token max lifetime (HDFS/Hive/HBase/Kafka/Ozone)", "15s", "Token expires for renewal tests"),
        ("Token renew interval", "8s", "Service-side renewal window"),
        ("Livy renewal interval", "5s", "livy.impersonation.delegation-token.renewal.interval"),
        ("Renewal job sleep", "45s", "spark.livy.test.renewal.sleep.seconds"),
        ("Renewal check interval", "8s", "spark.livy.test.renewal.check.interval.seconds"),
    ])

    add_heading(doc, "8.4. Auto-populated cluster properties", 2)
    add_table(doc, ["Key", "Source", "Example"], [
        ("hive.metastore.uris", "HiveMiniClusterSupport", "thrift://localhost:PORT"),
        ("hbase.zookeeper.quorum", "HBaseMiniClusterSupport", "localhost"),
        ("kafka.bootstrap.servers", "KafkaMiniClusterSupport", "127.0.0.1:PORT"),
        ("livy.test.ozone.path", "OzoneMiniClusterSupport", "ofs://omService/livy-it/bucket1"),
        ("hive.started / hbase.started / kafka.started", "Mini cluster support", "true"),
        ("ozone.started", "OzoneMiniClusterSupport", "true"),
    ])

    add_heading(doc, "8.5. External cluster (cluster.spec)", 2)
    add_para(
        doc,
        "In external mode (cluster.type=external), a real cluster is configured from "
        "cluster.spec.template. Instead of embedded mini-cluster services, tests use "
        "endpoints from the spec (hive.metastore.uris, hbase.zookeeper.quorum, "
        "kafka.bootstrap.servers).")

    add_code_block(doc, """# Example: external cluster
cluster.type=external
authScheme=kerberos
configDir=/etc/hadoop/conf
livyEndpoint=https://livy.example.com
principal=livy/host@REALM
keytabPath=/etc/security/keytabs/livy.keytab
delegation.token.services=hdfs,hive,hbase,kafka""")

    add_heading(doc, "8.6. Maven dependencies (test scope, parent pom.xml)", 2)
    add_table(doc, ["Artifact", "Version", "Purpose"], [
        ("ozone-mini-cluster, ozone-client, ozone-filesystem-hadoop3", "2.1.1", "MiniOzoneCluster"),
        ("hive-service", "3.0.0", "Hive Metastore Thrift server"),
        ("hbase-server (classifier: tests)", "2.4.17", "HBaseTestingUtility"),
        ("kafka_2.12", "2.8.2", "Embedded Kafka broker"),
        ("curator-test", "5.9.0", "ZK for Kafka broker"),
        ("hadoop-minikdc", "${hadoop.version}", "MiniKdc"),
    ])

    add_heading(doc, "9. Test cases", 1)

    add_heading(doc, "9.1. Unit tests (server module)", 2)
    add_table(doc, ["Class", "Coverage"], [
        ("DelegationTokenManagerSpec", "isEnabled, serialize/deserialize, RPC bootstrap class"),
        ("DelegationTokenProviderRegistrySpec", "Built-in providers, custom provider config, SPI"),
        ("DelegationTokenKerberosSpec", "MiniKdc, HBase/Kafka stub tokens, renewer startup"),
        ("SessionFilesystemUriCollectorSpec", "hdfs://, ofs://, o3fs:// URI extraction from session inputs"),
    ])

    add_heading(doc, "9.2. Integration tests – DelegationTokenIT matrix", 2)

    add_para(doc, "Dimensions:", bold=True)
    add_bullet(doc, "Deploy mode: client, cluster")
    add_bullet(doc, "Session type: interactive, batch")
    add_bullet(doc, "Impersonation: no-impersonation, impersonation (proxy-user, Kerberos only)")
    add_bullet(doc, "Services: hdfs, hive, hbase, kafka (+ ozone when started)")

    add_heading(doc, "9.2.1. Mini cluster regression (non-Kerberos HDFS)", 2)
    add_table(doc, ["Test ID", "Description"], [
        ("IT-DT-mini-interactive-{client|cluster}-no-impersonation-hdfs",
         "Interactive: no polling path, HDFS mkdir/exists"),
        ("IT-DT-mini-batch-{client|cluster}-no-impersonation-hdfs",
         "Batch: batch-dt-services.py succeeds"),
    ])

    add_heading(doc, "9.2.2. Kerberos service access", 2)
    add_table(doc, ["Test ID", "Description"], [
        ("IT-DT-krb-interactive-{client|cluster}-{no-impersonation|impersonation}-{service}",
         "Interactive: verifyNoFilesystemTokenPolling + verifyServiceAccessInteractive"),
        ("IT-DT-krb-batch-{client|cluster}-{no-impersonation|impersonation}-{service}",
         "Batch: batch-dt-services.py narrowed to one service"),
    ])
    add_para(doc, "service ∈ {hdfs, hive, hbase, kafka} – only when the embedded service started "
             "(hive.started, hbase.started, kafka.started).")

    add_heading(doc, "9.2.3. Delegation token renewal (short lifetime + long job)", 2)
    add_table(doc, ["Test ID", "Description"], [
        ("IT-DT-renewal-batch-{client|cluster}-impersonation-{service}",
         "Batch: batch-dt-renewal.py, 45s sleep, service access every 8s"),
        ("IT-DT-renewal-interactive-{client|cluster}-impersonation-hdfs",
         "Interactive: 45s HDFS mkdir loop, at least 2 successful checks"),
        ("IT-DT-renewal-interactive-{client|cluster}-impersonation-ozone",
         "Interactive: 45s ofs:// mkdir loop (when ozone.started=true)"),
    ])
    add_para(doc, "renewal batch service ∈ testServices + ozone (when configured).")

    add_heading(doc, "9.2.4. Ozone auto-discovery (ofs://)", 2)
    add_table(doc, ["Test ID", "Description"], [
        ("IT-DT-O01", "Interactive cluster: spark.files ofs:// path → Ozone mkdir/exists"),
        ("IT-DT-O02", "Batch cluster: auto-discovered ofs:// token + batch-dt-services.py"),
    ])

    add_heading(doc, "9.2.5. Legacy smoke tests", 2)
    add_table(doc, ["Test ID", "Description"], [
        ("IT-DT-01", "Interactive cluster: no filesystem token polling, 1+1"),
        ("IT-DT-02", "Batch cluster: batch-dt-test.py succeeds"),
        ("IT-DT-K01", "Interactive cluster + proxy: HDFS access"),
        ("IT-DT-K03", "Batch cluster + proxy: batch-dt-krb.py"),
    ])

    add_heading(doc, "9.3. Test scripts (integration-test resources)", 2)
    add_table(doc, ["File", "Role"], [
        ("batch-dt-services.py", "Service-specific access (hdfs/hive/hbase/kafka)"),
        ("batch-dt-renewal.py", "Long sleep + periodic service access for renewal IT"),
        ("batch-dt-krb.py", "Kerberos batch smoke"),
        ("batch-dt-test.py", "Basic batch smoke"),
    ])

    add_para(doc, "Verifications:", bold=True)
    add_bullet(doc, "verifyNoFilesystemTokenPolling: sc.hadoopConfiguration.get(\"livy.rsc.delegation.tokens.path\") == null")
    add_bullet(doc, "verifyHdfsAccessInteractive: fs.mkdirs + fs.exists")
    add_bullet(doc, "verifyHiveAccessInteractive: spark.sql(\"SHOW DATABASES\").count() > 0")
    add_bullet(doc, "verifyHbaseAccessInteractive: ConnectionFactory + admin.listNamespaceDescriptors")
    add_bullet(doc, "verifyKafkaAccessInteractive: KafkaConsumer bootstrap connect + close")
    add_bullet(doc, "verifyOzoneAccessInteractive: ofs:// path mkdir/exists")

    add_heading(doc, "10. Running tests", 1)

    add_code_block(doc, """# Build
mvn -pl server,integration-test -am -Pscala-2.12 -Pspark3 clean install \\
  -DskipTests -DskipITs -Drat.skip=true -Dcheckstyle.skip=true -Dscalastyle.skip=true

# Delegation token tests (unit + integration)
mvn -pl server,integration-test -Pscala-2.12 -Pspark3 test integration-test \\
  -DskipITs=false -Drat.skip=true -Dcheckstyle.skip=true \\
  -DwildcardSuites=org.apache.livy.utils.DelegationTokenManagerSpec,\\
org.apache.livy.utils.DelegationTokenProviderRegistrySpec,\\
org.apache.livy.utils.DelegationTokenKerberosSpec,\\
org.apache.livy.utils.SessionFilesystemUriCollectorSpec,\\
org.apache.livy.test.DelegationTokenIT

# Integration tests only (embedded mini cluster)
mvn -pl integration-test -Pscala-2.12 -Pspark3 integration-test \\
  -DskipITs=false -Drat.skip=true \\
  -DwildcardSuites=org.apache.livy.test.DelegationTokenIT

# Against external cluster (custom cluster.spec)
mvn -pl integration-test -Pscala-2.12 -Pspark3 integration-test \\
  -DskipITs=false -Dcluster.spec=/path/to/cluster.spec \\
  -DwildcardSuites=org.apache.livy.test.DelegationTokenIT""")

    add_heading(doc, "11. Log messages (troubleshooting)", 1)

    add_table(doc, ["Component", "Message", "Meaning"], [
        ("Livy", "Delegation token RPC registry listening on ...", "Batch RPC registry started"),
        ("Livy", "Batch session N connected for delegation token RPC", "Batch driver registered"),
        ("Livy", "Renewed delegation tokens for session N", "Successful renewal and push"),
        ("Driver", "Starting delegation token RPC server for batch driver", "DelegationTokenRpcBootstrap active"),
        ("Driver", "Applied delegation tokens from Livy RPC to driver UGI", "Tokens applied"),
    ])

    doc.save(OUTPUT)
    print(f"Generated: {OUTPUT}")


if __name__ == "__main__":
    build_document()
